#!/usr/bin/env python3
"""Hreinn Word-grunnur með íslenskum stílum; engin tenging við frumskjal."""
import argparse
from copy import deepcopy
from lxml import etree
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

DEFAULT_LAWYER = '[Nafn lögmanns]'


def element(tag, **attributes):
    node = OxmlElement('w:' + tag)
    for key, value in attributes.items():
        node.set(qn('w:' + key), str(value))
    return node


def enforce_monochrome(doc, font_name, body_size):
    """Fjarlægja litað Word-þema, líka úr tengdum og ónotuðum stílum."""
    for s in doc.styles:
        if s.type not in (WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER):
            continue
        s.font.name = font_name
        s.font.color.rgb = RGBColor(0, 0, 0)
        fonts = s.element.get_or_add_rPr().find(qn('w:rFonts'))
        fonts.attrib.clear()
        for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            fonts.set(qn('w:' + key), font_name)
        if s.name.lower().startswith(('heading ', 'subtitle')):
            s.font.size = Pt(body_size)
        if s.type == WD_STYLE_TYPE.PARAGRAPH and s.name.lower().startswith('heading '):
            s.paragraph_format.keep_with_next = True
    # Stafastíll tengdur fyrirsögn má ekki endurvekja blátt letur við ritun.
    for s in doc.styles:
        if s.type != WD_STYLE_TYPE.PARAGRAPH:
            continue
        link = s.element.find(qn('w:link'))
        if link is None or s.element.rPr is None:
            continue
        linked = next((x for x in doc.styles if x.style_id == link.get(qn('w:val'))), None)
        if linked is not None and linked.type == WD_STYLE_TYPE.CHARACTER:
            old = linked.element.find(qn('w:rPr'))
            if old is not None:
                linked.element.remove(old)
            linked.element.append(deepcopy(s.element.rPr))
    for node in doc.styles.element.xpath('.//w:color'):
        node.attrib.clear()
        node.set(qn('w:val'), '000000')
    for node in doc.styles.element.xpath('.//w:pBdr | .//w:shd'):
        node.getparent().remove(node)
    # Þemalitir geta annars birst aftur þegar Word býr til nýja stíla.
    a = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
    for rel in doc.part.rels.values():
        if rel.reltype.endswith('/theme'):
            part = rel.target_part
            theme = etree.fromstring(part.blob)
            for scheme in theme.iter(a + 'clrScheme'):
                for slot in scheme:
                    for child in list(slot):
                        slot.remove(child)
                    color = etree.SubElement(slot, a + 'srgbClr')
                    color.set('val', 'FFFFFF' if etree.QName(slot).localname.startswith('lt') else '000000')
            for latin in theme.iter(a + 'latin'):
                latin.set('typeface', font_name)
            part._blob = etree.tostring(theme, xml_declaration=True, encoding='UTF-8', standalone=True)


def add_firm_header(header, text='Merki Stofu'):
    """Setja breytanlegan texta í skjalahaus, án merkismyndar."""
    p = header.paragraphs[0]
    p.clear()
    p.style = 'Header'
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1
    p.add_run(text)


def style(doc, name, size=10, bold=False, italic=False, underline=False,
          alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, line=259 / 240):
    styles = doc.styles
    s = styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    if name != 'Normal':
        s.base_style = styles['Normal']
    s.font.name = 'Verdana'
    s.font.size = Pt(size)
    s.font.bold, s.font.italic, s.font.underline = bold, italic, underline
    s.font.color.rgb = RGBColor(0, 0, 0)
    rpr = s.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    for attr in list(fonts.attrib):
        if 'theme' in attr.lower():
            del fonts.attrib[attr]
    for name_ in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + name_), 'Verdana')
    rpr.append(element('lang', val='is-IS'))
    fmt = s.paragraph_format
    fmt.alignment = alignment
    fmt.space_before, fmt.space_after = Pt(0), Pt(8)
    fmt.line_spacing = line
    fmt.left_indent = fmt.right_indent = fmt.first_line_indent = Mm(0)
    fmt.widow_control = True
    ppr = s.element.get_or_add_pPr()
    for border in list(ppr.findall(qn('w:pBdr'))):
        ppr.remove(border)
    s.next_paragraph_style = styles['Normal']
    return s


def numbering(doc, fmt, bold=False, italic=False, underline=False):
    root = doc.part.numbering_part.element
    aid = max([int(n.get(qn('w:abstractNumId'))) for n in root.findall(qn('w:abstractNum'))] + [-1]) + 1
    nid = max([int(n.get(qn('w:numId'))) for n in root.findall(qn('w:num'))] + [0]) + 1
    abstract = element('abstractNum', abstractNumId=aid)
    abstract.append(element('multiLevelType', val='singleLevel'))
    level = element('lvl', ilvl=0)
    for item in (element('start', val=1), element('numFmt', val=fmt),
                 element('lvlText', val='%1.'), element('suff', val='tab'),
                 element('lvlJc', val='left')):
        level.append(item)
    ppr = element('pPr')
    tabs = element('tabs')
    tabs.append(element('tab', val='num', pos=1080))
    ppr.append(tabs)
    ppr.append(element('ind', left=1080, hanging=720))
    level.append(ppr)
    rpr = element('rPr')
    rpr.append(element('rFonts', ascii='Verdana', hAnsi='Verdana'))
    rpr.append(element('sz', val=20))
    rpr.append(element('b', val=int(bold)))
    rpr.append(element('i', val=int(italic)))
    rpr.append(element('u', val='none'))
    level.append(rpr)
    abstract.append(level)
    first_num = root.find(qn('w:num'))
    if first_num is None:
        root.append(abstract)
    else:
        first_num.addprevious(abstract)
    num = element('num', numId=nid)
    num.append(element('abstractNumId', val=aid))
    root.append(num)
    return nid


def bind_number(s, nid):
    ppr = s.element.get_or_add_pPr()
    old = ppr.find(qn('w:numPr'))
    if old is not None:
        ppr.remove(old)
    numpr = element('numPr')
    numpr.append(element('ilvl', val=0))
    numpr.append(element('numId', val=nid))
    ppr.append(numpr)
    s.paragraph_format.left_indent = Mm(19.05)
    s.paragraph_format.first_line_indent = Mm(-12.7)


def create_document(header_text='Merki Stofu'):
    doc = Document()
    # Hreinsa óþarfa aukahluta úr sjálfgefnum python-docx grunni.
    for owner in (doc.part, doc.part.package):
        for rid, rel in list(owner.rels.items()):
            if rel.reltype.endswith(('/customXml', '/stylesWithEffects', '/thumbnail')):
                owner.drop_rel(rid) if hasattr(owner, 'drop_rel') else owner.rels.pop(rid)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Mm(25.4)
    sec.header_distance = sec.footer_distance = Mm(12.49)
    sec.different_first_page_header_footer = True
    # OOXML-röð skiptir máli: síðutal verður að koma á undan dálkum og línuristi.
    sec._sectPr.insert_element_before(element('pgNumType', start=1), 'w:cols', 'w:formProt', 'w:vAlign', 'w:noEndnote', 'w:titlePg', 'w:textDirection', 'w:bidi', 'w:rtlGutter', 'w:docGrid')
    style(doc, 'Normal')
    title = style(doc, 'Title', size=12, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    title.element.get_or_add_rPr().append(element('spacing', val=60))
    title.paragraph_format.keep_with_next = True
    for name, bold, italic, underline, fmt in (
        ('Heading 1', True, False, False, 'upperRoman'),
        ('Heading 2', False, True, True, 'lowerRoman'),
    ):
        s = style(doc, name, bold=bold, italic=italic, underline=underline,
                  alignment=WD_ALIGN_PARAGRAPH.LEFT, line=1.15 if bold else 259 / 240)
        s.paragraph_format.keep_with_next = s.paragraph_format.keep_together = True
        bind_number(s, numbering(doc, fmt, bold, italic, underline))
    s = style(doc, 'Logmadur Viðtakandi', alignment=WD_ALIGN_PARAGRAPH.LEFT)
    s.paragraph_format.space_before, s.paragraph_format.space_after = Pt(14), Pt(28)
    s = style(doc, 'Logmadur Dagsetning', alignment=WD_ALIGN_PARAGRAPH.RIGHT)
    s.paragraph_format.space_after = Pt(28)
    s.paragraph_format.keep_with_next = True
    style(doc, 'Logmadur Inngangur', line=1.15)
    s = style(doc, 'Logmadur Undirritun', alignment=WD_ALIGN_PARAGRAPH.CENTER)
    s.paragraph_format.keep_with_next = s.paragraph_format.keep_together = True
    style(doc, 'Logmadur Fylgiskjal', italic=True, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    for name in ('Header', 'Footer'):
        s = style(doc, name, alignment=WD_ALIGN_PARAGRAPH.CENTER, line=1)
        s.paragraph_format.space_after = Pt(0)
    sec.first_page_header.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for footer in (sec.footer, sec.first_page_footer):
        p = footer.paragraphs[0]
        p.style = doc.styles['Footer']
        field = element('fldSimple', instr=' PAGE ')
        run = element('r')
        text = element('t')
        text.text = '1'
        run.append(text)
        field.append(run)
        p._p.append(field)
    doc.settings.element.append(element('updateFields', val='true'))
    props = doc.core_properties
    for key in ('author', 'last_modified_by', 'title', 'subject', 'keywords',
                'comments', 'category', 'identifier', 'content_status'):
        setattr(props, key, '')
    props.language = 'is-IS'
    add_firm_header(sec.first_page_header, header_text)
    enforce_monochrome(doc, 'Verdana', 10)
    return doc


def add_title(doc, text):
    return doc.add_paragraph(text.upper(), 'Title')


def add_heading(doc, text, level=1, restart=False):
    if level not in (1, 2):
        raise ValueError('Aðeins fyrirsagnarstig 1 og 2 eru skilgreind.')
    name = 'Heading ' + str(level)
    if restart:
        # Fyrri málsgreinar halda númeraröð sinni með beinni numPr-stillingu.
        nid = numbering(doc, 'upperRoman' if level == 1 else 'lowerRoman',
                        bold=level == 1, italic=level == 2, underline=level == 2)
        bind_number(doc.styles[name], nid)
    p = doc.add_paragraph(text, name)
    nid = doc.styles[name].element.get_or_add_pPr().find(qn('w:numPr')).find(qn('w:numId')).get(qn('w:val'))
    numpr = element('numPr')
    numpr.append(element('ilvl', val=0))
    numpr.append(element('numId', val=nid))
    p._p.get_or_add_pPr().append(numpr)
    return p


def add_attachment(doc, number, title):
    if not isinstance(number, int) or isinstance(number, bool) or number < 1:
        raise ValueError('Fylgiskjalsnúmer verður að vera jákvæð heiltala.')
    if not title.strip():
        raise ValueError('Heiti fylgiskjals má ekki vera tómt.')
    return doc.add_paragraph(f'Fskj. {number}. {title}', 'Logmadur Fylgiskjal')


def save_document(doc, output):
    output = Path(output)
    if output.suffix.lower() != '.docx':
        raise ValueError('Úttaksskrá verður að enda á .docx.')
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open('xb') as stream:
        doc.save(stream)


def add_signature(doc, place_date, lawyer=DEFAULT_LAWYER, *, on_behalf_of=None):
    """Setja undirritunarreit fyrir tilgreindan lögmann, án rafrænnar undirskriftar."""
    lines = [place_date, 'Virðingarfyllst,']
    if on_behalf_of:
        lines.append(f'F.h. {on_behalf_of}')
    lines.append(f'{lawyer}, lögmaður')
    for text in lines:
        p = doc.add_paragraph(text, 'Logmadur Undirritun')
    p.paragraph_format.keep_with_next = False
    return p


def sample(doc):
    doc.add_paragraph('[Viðtakandi]', 'Logmadur Viðtakandi')
    doc.add_paragraph('[Staður], [dagsetning]', 'Logmadur Dagsetning')
    add_title(doc, 'Minnisblað')
    doc.add_paragraph('Þetta er hlutlaust sniðpróf. Textinn sýnir framsetningu og hefur ekkert efnislegt gildi.', 'Logmadur Inngangur')
    add_heading(doc, 'Tilgangur')
    doc.add_paragraph('Tilgangur þessa skjals er að sýna samræmi í letri, línubili og kaflaskiptingu. Nýtt skjal fær sjálfstætt efni samkvæmt fyrirmælum notanda.')
    add_heading(doc, 'Forsendur og mat')
    add_heading(doc, 'Gögn og afmörkun', level=2)
    doc.add_paragraph('Í nýju verkefni þarf að greina hvaða upplýsingar liggja fyrir og hvaða atriði þarfnast nánari skýringar. Tilvísanir tengja umfjöllun við rétt gögn.')
    add_heading(doc, 'Atriði sem þarfnast nánari skoðunar og skýringar áður en unnt er að leggja endanlegt mat á fyrirliggjandi upplýsingar', level=2)
    doc.add_paragraph('Með hliðsjón af framangreindu skal rökstuðningur vera skýr og studdur gögnum. Ekki verður dregin efnisleg niðurstaða af þessu sniðprófi.')
    doc.add_page_break()
    add_heading(doc, 'Niðurstaða')
    add_heading(doc, 'Framhald', level=2, restart=True)
    doc.add_paragraph('Hér sést framhald skjals á annarri síðu. Blaðsíðutal á að uppfærast og stór rómversk kaflanúmer halda áfram í réttri röð.')
    add_signature(doc, '[Staður og dagsetning]')
    add_attachment(doc, 1, '[Heiti fyrsta fylgiskjals]')
    add_attachment(doc, 2, '[Heiti annars fylgiskjals]')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', required=True, type=Path, help='Slóð að nýrri DOCX-skrá.')
    parser.add_argument('--sample', action='store_true', help='Búa til hlutlaust sniðpróf.')
    args = parser.parse_args()
    doc = create_document()
    if args.sample:
        sample(doc)
    save_document(doc, args.output)
    print(args.output.resolve())


if __name__ == '__main__':
    main()
